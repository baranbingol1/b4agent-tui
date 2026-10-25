#!/usr/bin/env python3
"""Recalculate a workbook's formulas headlessly with LibreOffice.

openpyxl never computes formulas. This script shells out to `soffice`
(LibreOffice) to open the workbook, recalculate, and re-save it, so
cached formula results become available to `xlsx_read.py --data-only`
and `--formulas`.

Behavior:
  * soffice on PATH: converts the file to .xlsx in a temp dir (which
    recalculates all formulas) and replaces the original (or writes
    --out). Prints {"recalculated": true, ...} and exits 0.
  * soffice absent: prints {"recalculated": false, "reason": ...} with
    installation guidance and STILL exits 0 — callers can branch on the
    JSON instead of the exit code.

Note: LibreOffice recalculates .xlsx on load per its default
calculation settings; conversion re-saves with fresh cached values.

Error scan (added by b4agent): every run also lists formula errors, so
a recalculated workbook is never reported as fine while it shows #REF!
or #NAME?. It reports cells whose cached value is an Excel error
(#REF!, #NAME?, #DIV/0!, #VALUE!, #N/A, #NUM!, #NULL!, ...) and
formulas whose text contains #REF! (a reference that was deleted). The
JSON gains "error_total", "error_counts" and "errors" (first 200 cells:
sheet, cell, error, formula). Without soffice the scan runs on the
cached values already in the file, which a file fresh from openpyxl
doesn't have. --strict exits 2 when any error is found.

Usage:
  xlsx_recalc.py book.xlsx
  xlsx_recalc.py book.xlsx --out recalced.xlsx
  xlsx_recalc.py book.xlsx --timeout 120
  xlsx_recalc.py book.xlsx --strict
"""
from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path


ERROR_VALUES = ("#REF!", "#NAME?", "#DIV/0!", "#VALUE!", "#N/A", "#NUM!",
                "#NULL!", "#SPILL!", "#CALC!", "#GETTING_DATA")
MAX_LISTED_ERRORS = 200


def scan(path):
    """Count formula cells and cached values, and list formula errors."""
    from openpyxl import load_workbook
    wb_f = load_workbook(path, data_only=False)
    wb_v = load_workbook(path, data_only=True)
    formulas = cached = 0
    errors, counts = [], {}
    for name in wb_f.sheetnames:
        ws_f, ws_v = wb_f[name], wb_v[name]
        for row in ws_f.iter_rows():
            for cell in row:
                raw = cell.value
                text = getattr(raw, "text", raw)  # ArrayFormula keeps it in .text
                formula = text if isinstance(text, str) and text.startswith("=") else None
                value = ws_v[cell.coordinate].value
                if formula:
                    formulas += 1
                    if value is not None:
                        cached += 1
                # A deleted reference in the text is the root cause even when
                # the cached value says otherwise (LibreOffice caches
                # =SUM(#REF!) as #NAME?).
                err = None
                if formula and "#REF!" in formula.upper():
                    err = "#REF!"
                elif isinstance(value, str) and value.strip() in ERROR_VALUES:
                    err = value.strip()
                if err:
                    counts[err] = counts.get(err, 0) + 1
                    if len(errors) < MAX_LISTED_ERRORS:
                        errors.append({"sheet": name, "cell": cell.coordinate,
                                       "error": err, "formula": formula})
    return formulas, cached, {"error_total": sum(counts.values()),
                              "error_counts": counts, "errors": errors}


def main(argv=None):
    ap = argparse.ArgumentParser(
        description="Recalculate .xlsx formulas headlessly via LibreOffice.")
    ap.add_argument("file", help="path to .xlsx file")
    ap.add_argument("--out", help="output path (default: replace input)")
    ap.add_argument("--timeout", type=int, default=180,
                    help="seconds to wait for soffice (default 180)")
    ap.add_argument("--strict", action="store_true",
                    help="exit 2 when the scan finds formula errors")
    args = ap.parse_args(argv)

    src = Path(args.file).resolve()
    if not src.exists():
        print(json.dumps({"ok": False, "error": f"no such file: {src}"}),
              file=sys.stderr)
        return 1

    soffice = shutil.which("soffice")
    if not soffice:
        formulas, cached, found = scan(src)
        print(json.dumps({
            "ok": True, "recalculated": False,
            "reason": "LibreOffice (soffice) not found on PATH",
            "guidance": "Install LibreOffice (e.g. `apt install "
                        "libreoffice-calc` or `brew install --cask "
                        "libreoffice`), or open the file in Excel/"
                        "LibreOffice once and re-save it.",
            "formula_cells": formulas, "with_cached_values": cached, **found,
        }, ensure_ascii=False))
        return 2 if args.strict and found["error_total"] else 0

    with tempfile.TemporaryDirectory() as tmp:
        proc = subprocess.run(
            [soffice, "--headless", "--calc", "--convert-to", "xlsx:Calc "
             "MS Excel 2007 XML", "--outdir", tmp, str(src)],
            capture_output=True, text=True, encoding="utf-8",
            timeout=args.timeout,
            env={"HOME": tmp, "PATH": Path(soffice).parent.as_posix()
                 + ":/usr/bin:/bin"}, check=False)
        produced = Path(tmp) / (src.stem + ".xlsx")
        if proc.returncode != 0 or not produced.exists():
            print(json.dumps({"ok": False,
                              "error": "soffice conversion failed",
                              "stderr": proc.stderr.strip()[-500:]}),
                  file=sys.stderr)
            return 1
        formulas, cached, found = scan(produced)
        dest = Path(args.out).resolve() if args.out else src
        shutil.copyfile(produced, dest)

    print(json.dumps({
        "ok": True, "recalculated": True, "output": str(dest),
        "formula_cells": formulas, "with_cached_values": cached, **found,
    }, ensure_ascii=False))
    return 2 if args.strict and found["error_total"] else 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}), file=sys.stderr)
        sys.exit(1)
